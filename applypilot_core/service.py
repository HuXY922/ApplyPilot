"""Application use cases. Runtime calls are lazy, bounded by their source adapters.

No use case performs external applications, uploads, authentication or form fills.
"""
import base64
import binascii
from copy import deepcopy
import importlib
import io
import json
import re
import subprocess
import sys
import threading
import zipfile

from .store import DomainError, Store, encode, require

TRACK_STATUSES = {'saved','skipped','preparing','filling','waiting_user','submission_unknown','submitted','assessment','interview','offer','rejected'}
PROFILE_FIELDS = {'name','education','experience','projects','skills','preferences','resume_text'}
MAX_IMPORT_BYTES = 2*1024*1024
MODEL_PROVIDERS = ('openai','deepseek')
# The supported deployment uses one Gunicorn process. Do not queue parser work:
# at most one 256 MiB-bounded subprocess may run in that process at a time.
_IMPORT_SLOT = threading.BoundedSemaphore(1)


def string(value, field, limit=2000, allow_empty=True):
    require(isinstance(value,str), 'invalid_input', field+' must be text.')
    require(len(value) <= limit and not re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',value), 'invalid_input', field+' is too long or contains control characters.')
    require(allow_empty or bool(value.strip()), 'invalid_input', field+' is required.')
    return value.strip()


def string_list(value, field, count=50, size=200):
    require(isinstance(value,list) and len(value) <= count, 'invalid_input', field+' must be a bounded list.')
    return list(dict.fromkeys(string(v,field,size,False) for v in value))


def validate_profile(profile):
    require(isinstance(profile,dict) and set(profile) <= PROFILE_FIELDS, 'invalid_profile', 'Profile fields are invalid.')
    result = {'name': string(profile.get('name',''),'name',200),
              'resume_text': string(profile.get('resume_text',''),'resume_text',30000)}
    specs = {'education': {'institution':200,'degree':100,'major':200,'graduation':20},
             'experience': {'id':100,'organization':200,'role':200,'description':6000},
             'projects': {'id':100,'name':200,'description':6000}}
    for field, fields in specs.items():
        values = profile.get(field,[])
        require(isinstance(values,list) and len(values) <= 20, 'invalid_profile', field+' must contain at most 20 entries.')
        result[field] = []
        ids = set()
        for i, value in enumerate(values):
            require(isinstance(value,dict) and set(value) <= set(fields), 'invalid_profile', field+' entry contains unsupported fields.')
            item = {k:string(value.get(k,''),field+'.'+k,limit) for k,limit in fields.items()}
            if 'id' in item:
                item['id'] = item['id'] or field+'-'+str(i+1)
                require(item['id'] not in ids, 'invalid_profile', 'Evidence identifiers must be unique.')
                ids.add(item['id'])
            if field == 'education' and item['graduation']:
                require(re.fullmatch(r'20\d{2}(?:-(?:0[1-9]|1[0-2]))?',item['graduation']),
                        'invalid_profile', 'Graduation must be YYYY or YYYY-MM.')
            result[field].append(item)
    result['skills'] = string_list(profile.get('skills',[]),'skills')
    prefs = profile.get('preferences',{})
    require(isinstance(prefs,dict) and set(prefs) <= {'roles','cities','excluded_roles','recruitment_type','graduation_year'},
            'invalid_profile','Preferences contain unsupported fields.')
    result['preferences'] = {k:string_list(prefs.get(k,[]),k,20,100) for k in ('roles','cities','excluded_roles')}
    recruitment = prefs.get('recruitment_type','campus')
    require(recruitment in ('campus','internship','all'),'invalid_input','Unknown recruitment type.')
    result['preferences']['recruitment_type'] = recruitment
    year = string(prefs.get('graduation_year',''),'graduation_year',4)
    require(not year or bool(re.fullmatch('20[0-9]{2}',year)), 'invalid_profile', 'Graduation year must be YYYY.')
    result['preferences']['graduation_year'] = year
    require(len(encode(result)) <= 80000, 'invalid_profile', 'Profile exceeds the demo size limit.')
    return result


# Running document parsing in a separate process gives both time and memory caps.
# Input remains local and is never written to a user-specified path or sent away.
_PARSER = r'''
import io,json,sys
try:
    import resource
    resource.setrlimit(resource.RLIMIT_AS,(256*1024*1024,256*1024*1024))
    resource.setrlimit(resource.RLIMIT_CPU,(8,8))
except ImportError:
    pass
raw=sys.stdin.buffer.read(2*1024*1024+1)
ext=sys.argv[1]
if ext=="pdf":
    from pypdf import PdfReader
    import pypdf.filters
    if hasattr(pypdf.filters,"ZLIB_MAX_OUTPUT_LENGTH"):
        pypdf.filters.ZLIB_MAX_OUTPUT_LENGTH=8*1024*1024
    reader=PdfReader(io.BytesIO(raw),strict=True)
    if reader.is_encrypted: raise ValueError("encrypted")
    if len(reader.pages)>30: raise ValueError("page_limit")
    texts=[]
    for page in reader.pages:
        text=page.extract_text() or ""
        if len(text)>30000 or sum(map(len,texts))+len(text)>30000: raise ValueError("text_limit")
        texts.append(text)
    text="\n".join(texts)
else:
    from docx import Document
    doc=Document(io.BytesIO(raw))
    text="\n".join([p.text for p in doc.paragraphs]+[" | ".join(c.text for c in r.cells) for t in doc.tables for r in t.rows])
    if len(text)>30000: raise ValueError("text_limit")
print(json.dumps({"text":text},ensure_ascii=False))
'''


def import_text(payload):
    require(isinstance(payload,dict), 'invalid_input', 'Expected a JSON object.')
    if 'text' in payload:
        require(set(payload) == {'text'}, 'invalid_import', 'Choose text or one file, not both.')
        return string(payload['text'],'text',30000,False), 'text'
    require(set(payload) <= {'filename','content_base64','mime_type'} and {'filename','content_base64'} <= set(payload),
            'invalid_import', 'Import plain text, TXT, PDF or DOCX.')
    filename = string(payload['filename'],'filename',200,False)
    require('/' not in filename and '\\' not in filename and filename not in ('.','..'), 'invalid_import', 'Invalid filename.')
    ext = filename.rsplit('.',1)[-1].lower()
    require(ext in ('pdf','docx','txt'), 'unsupported_file', 'Only TXT, PDF and DOCX files are supported.', 415)
    mime = payload.get('mime_type')
    allowed_mime = {'pdf':'application/pdf','docx':'application/vnd.openxmlformats-officedocument.wordprocessingml.document','txt':'text/plain'}
    require(mime is None or mime == allowed_mime[ext], 'unsupported_file', 'File MIME type does not match its extension.',415)
    content = payload['content_base64']
    require(isinstance(content,str) and len(content) <= (MAX_IMPORT_BYTES+2)//3*4, 'file_too_large', 'File must be at most 2 MiB.', 413)
    try:
        raw = base64.b64decode(content,validate=True)
    except (ValueError,binascii.Error):
        raise DomainError(400,'invalid_import','File encoding is invalid.') from None
    require(0 < len(raw) <= MAX_IMPORT_BYTES, 'file_too_large', 'File must be at most 2 MiB.', 413)
    if ext == 'txt':
        try:
            return string(raw.decode('utf-8-sig'),'text',30000,False), ext
        except UnicodeDecodeError:
            raise DomainError(400,'invalid_import','TXT must use UTF-8 encoding.') from None
    if ext == 'pdf':
        require(raw.startswith(b'%PDF-'), 'invalid_import', 'This is not a supported PDF file.')
    if ext == 'docx':
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                entries = archive.infolist()
                require(len(entries) <= 1000 and sum(e.file_size for e in entries) <= 8*1024*1024,
                        'unsafe_document', 'DOCX expanded contents exceed the safety limit.')
                names = [e.filename for e in entries]
                require('[Content_Types].xml' in names and 'word/document.xml' in names,
                        'invalid_import', 'This is not a valid DOCX document.')
                for entry in entries:
                    require(not entry.filename.startswith(('/','\\')) and '..' not in entry.filename.replace('\\','/').split('/') and not entry.flag_bits & 1,
                            'unsafe_document', 'Unsafe or encrypted DOCX contents.')
                    require(entry.file_size <= 4*1024*1024 and entry.file_size <= max(entry.compress_size,1)*200,
                            'unsafe_document', 'DOCX compressed contents exceed the safety limit.')
                    require('vbaproject' not in entry.filename.lower(), 'unsafe_document', 'Macro-enabled documents are unsupported.')
                    if entry.filename.lower().endswith(('.xml','.rels')):
                        xml = archive.read(entry)
                        require(b'<!DOCTYPE' not in xml.upper() and b'<!ENTITY' not in xml.upper(), 'unsafe_document', 'Embedded XML declarations are unsupported.')
        except (zipfile.BadZipFile,RuntimeError,NotImplementedError,ValueError):
            raise DomainError(400,'invalid_import','DOCX could not be safely read.') from None
    try:
        parsed = subprocess.run([sys.executable,'-c',_PARSER,ext],input=raw,stdout=subprocess.PIPE,
                                stderr=subprocess.DEVNULL,timeout=12,check=False)
        require(parsed.returncode == 0 and len(parsed.stdout) <= 150000, 'document_parse_failed',
                'File could not be safely extracted. Use an unencrypted PDF (up to 30 pages), DOCX, or paste text.')
        text = json.loads(parsed.stdout)['text']
    except (subprocess.TimeoutExpired,ValueError,KeyError,OSError):
        raise DomainError(400,'document_parse_failed','File extraction failed or exceeded its safety limit. Try pasting text.') from None
    require(isinstance(text,str) and bool(text.strip()), 'no_text', 'No readable text found. Scanned PDFs need OCR first; paste extracted text.')
    return string(text,'extracted text',30000,False), ext


class Service:
    def __init__(self, store, runtime=None, reserve=None):
        self.store = store
        self._runtime = runtime
        self.reserve = reserve or (lambda kind: None)

    @property
    def runtime(self):
        if self._runtime is None:
            self._runtime = importlib.import_module('applypilot_core.runtime')
        return self._runtime

    def state(self):
        result = self.store.state()
        status = getattr(self.runtime,'model_status',None)
        result['model'] = status() if status else {'configured':False,'mode':'rules','status':'not_configured'}
        result['demo'] = {'public_demo':True,'external_submission_enabled':False,'fictional_data_only':True,
                          'notice':'Public demo: use fictional data only. Live searches send job filters, never your profile. No external applications are submitted.'}
        return result

    def import_profile(self, data):
        require(_IMPORT_SLOT.acquire(blocking=False), 'import_busy',
                'Another resume import is running. Please try again shortly.',429)
        try:
            return self._import_profile(data)
        finally:
            _IMPORT_SLOT.release()

    def _import_profile(self, data):
        text, file_type = import_text(data)
        try:
            draft = validate_profile(self.runtime.extract_profile(text))
        except DomainError:
            raise
        except Exception:
            raise DomainError(422,'extraction_failed','Resume extraction failed. Try shorter plain text.') from None
        self.store.save_draft(draft)
        return {'draft':draft,'requires_review':True,'mode':'rules','file_type':file_type}

    def save_profile(self, data):
        require(set(data) <= {'profile','confirmed','expected_version'}, 'invalid_input', 'Unexpected profile request fields.')
        require(data.get('confirmed') is True, 'confirmation_required', 'Review and confirm the profile before saving.', 409)
        require(type(data.get('expected_version')) is int, 'invalid_input', 'expected_version is required.')
        return self.store.save_profile(validate_profile(data.get('profile')),data['expected_version'])

    def search(self, data):
        require(set(data) <= {'query','recruitment_type','cities'}, 'invalid_input', 'Unexpected search fields. Source URLs cannot be supplied.')
        query = string(data.get('query',''),'query',200)
        require(not re.search(r'https?://|[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}|(?<!\d)\d{11,}(?!\d)',query),
                'invalid_search', 'Search using job terms only, without URLs or personal contact details.')
        recruitment = data.get('recruitment_type','campus')
        require(recruitment in ('campus','internship','all'), 'invalid_input', 'Unknown recruitment type.')
        cities = string_list(data.get('cities',[]),'cities',10,50)
        request = {'query':query,'recruitment_type':recruitment,'cities':cities}
        self.reserve('search')
        rid = self.store.begin_run('search',request)
        try:
            result = self.runtime.discover(**request)
            require(isinstance(result,dict) and isinstance(result.get('jobs'),list) and len(result['jobs']) <= 100,
                    'source_schema','Discovery returned an invalid result.',502)
            failures = result.get('failures',[])
            summaries = result.get('source_summary',[])
            require(isinstance(failures,list) and isinstance(summaries,list), 'source_schema', 'Discovery summary is invalid.',502)
            observations = result.get('observed_jobs',[])
            require(isinstance(observations,list) and len(observations) <= 1500,'source_schema','Invalid observation list.',502)
            refreshed = self.store.refresh_observed(observations,rid)
            if failures and not result['jobs']:
                run = self.store.finish_run(rid,'failed',{'failures':failures,'source_summary':summaries,'found':0,'created':0})
                raise DomainError(502,'source_unavailable','Live sources could not return verified results. Demo jobs were not substituted.',run=run)
            jobs, created = self.store.ingest(result['jobs'],rid)
            run = self.store.finish_run(rid,'partial' if failures else 'completed',
                                         {'failures':failures,'source_summary':summaries,'found':len(jobs),'created':created,'refreshed':refreshed})
            return {'jobs':jobs,'run':run}
        except DomainError as error:
            if not error.details.get('run'):
                self.store.finish_run(rid,'failed',{'code':error.code,'message':error.message})
            raise
        except Exception:
            run = self.store.finish_run(rid,'failed',{'code':'source_unavailable','message':'The live source request failed.'})
            raise DomainError(502,'source_unavailable','Live search failed. No demo results were substituted.',run=run) from None

    def generate(self, job_id, kind, data):
        require(set(data) <= {'model_consent'}, 'invalid_input', 'This action accepts only explicit model-provider consent.')
        state = self.store.profile()
        require(state['profile_confirmed'], 'profile_unconfirmed', 'Review and confirm your profile before analysis.', 409)
        job = self.store.job(job_id)
        model_status = getattr(self.runtime,'model_status',lambda: {'configured':False})()
        if model_status.get('configured'):
            provider = model_status.get('provider')
            require(provider in MODEL_PROVIDERS and data.get('model_consent') == provider,
                    'model_consent_required',
                    'Model-provider consent is missing or no longer matches current settings. Refresh and review before continuing.',409)
            self.reserve('model')
        elif 'model_consent' in data:
            require(data['model_consent'] in MODEL_PROVIDERS, 'invalid_input', 'Unknown model provider consent.')
        rid = self.store.begin_run(kind,{'job_id':job_id,'profile_version':state['profile_version']})
        try:
            func = self.runtime.analyze if kind == 'analysis' else self.runtime.suggest_resume
            result = func(deepcopy(state['profile']),deepcopy(job))
            require(isinstance(result,dict) and len(encode(result)) <= 100000, 'runtime_schema','Analysis returned an invalid result.',502)
            if kind == 'analysis':
                require(result.get('eligibility') in ('pass','fail','unknown'), 'runtime_schema','Analysis eligibility is missing.',502)
            else:
                require(isinstance(result.get('suggestions'),list), 'runtime_schema','Resume suggestions are missing.',502)
            degraded = str(result.get('model_status','')).startswith('fallback_rules:')
            if degraded:
                result = {**result,'degraded':True,'warning':'Model execution failed; showing explicitly labeled rule-based output.'}
            output = self.store.save_output(job_id,kind,state['profile_version'],job['observation_revision'],result)
            run = self.store.finish_run(rid,'partial' if degraded else 'completed',
                                        {'output_id':output['id'],'mode':output.get('mode','rules'),
                                         'model_status':output.get('model_status','unknown'),'degraded':degraded})
            return {kind:output,'run':run}
        except DomainError as error:
            self.store.finish_run(rid,'failed',{'code':error.code,'message':error.message})
            raise
        except Exception:
            self.store.finish_run(rid,'failed',{'code':'analysis_failed','message':'Analysis failed.'})
            raise DomainError(502,'analysis_failed','Analysis failed. Your profile and tracker were not changed.') from None

    def track(self, job_id, data):
        require(set(data) <= {'status','attestation','note','idempotency_key'}, 'invalid_input','Unexpected tracking fields.')
        status = data.get('status')
        require(isinstance(status,str) and status in TRACK_STATUSES, 'invalid_status','Unknown tracker status.')
        require('attestation' not in data or type(data['attestation']) is bool, 'invalid_input','attestation must be true or false.')
        key = string(data.get('idempotency_key'),'idempotency_key',160,False)
        require(bool(re.fullmatch(r'[A-Za-z0-9_.:-]{8,160}',key)), 'invalid_input', 'Use an action key of 8–160 safe characters.')
        note = string(data.get('note',''),'note',2000)
        return self.store.track(job_id,status,data.get('attestation',False),note,key)
