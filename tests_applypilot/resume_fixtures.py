"""Synthetic import fixtures: real TXT/PDF/DOCX bytes; never candidate data.

The selectable-text Chinese PDF was generated for this regression with ReportLab;
embedding its small output avoids adding a PDF-generation dependency. DOCX uses
our existing parser dependency and fixed metadata/ZIP dates for reproducible bytes.
"""
import base64
from datetime import datetime, timezone
import io
import zipfile

FICTIONAL_TEXT = '陈星河（虚构文件测试）\n毕业年份：2027\n教育经历\n虚构示例大学  硕士  信息管理  2027-07\n目标岗位：AI产品经理\n意向城市：上海，北京\n技能：Python，SQL，原型设计，用户研究\n项目经历\n校园学习助手（虚构教学项目）\n整理虚构教学小组的会议记录与任务清单。\n设计大模型Agent离线评测用例集，整理模型评测清单。\n使用Python汇总失败案例，根据模拟用户反馈调整交互原型。\n以上均为虚构教学资料，没有真实联系方式，不用于真实投递。\n'

_PDF_BASE64 = (
    'JVBERi0xLjMKJZOMi54gUmVwb3J0TGFiIEdlbmVyYXRlZCBQREYgZG9jdW1lbnQgKG9wZW5zb3VyY2UpCjEgMCBvYmoKPDwKL0Yx'
    'IDIgMCBSIC9GMiAzIDAgUgo+PgplbmRvYmoKMiAwIG9iago8PAovQmFzZUZvbnQgL0hlbHZldGljYSAvRW5jb2RpbmcgL1dpbkFu'
    'c2lFbmNvZGluZyAvTmFtZSAvRjEgL1N1YnR5cGUgL1R5cGUxIC9UeXBlIC9Gb250Cj4+CmVuZG9iagozIDAgb2JqCjw8Ci9CYXNl'
    'Rm9udCAvU1RTb25nLUxpZ2h0IC9EZXNjZW5kYW50Rm9udHMgWyA8PAovQmFzZUZvbnQgL1NUU29uZy1MaWdodCAvQ0lEU3lzdGVt'
    'SW5mbyA8PAovT3JkZXJpbmcgKEdCMSkgL1JlZ2lzdHJ5IChBZG9iZSkgL1N1cHBsZW1lbnQgMAo+PiAvRFcgMTAwMCAvRm9udERl'
    'c2NyaXB0b3IgPDwKL0FzY2VudCA3NTIgL0NhcEhlaWdodCA3MzcgL0Rlc2NlbnQgLTI3MSAvRmxhZ3MgNiAvRm9udEJCb3ggWyAt'
    'MjUgLTI1NCAxMDAwIDg4MCBdIC9Gb250TmFtZSAvU1RTb25nU3RkLUxpZ2h0IAogIC9JdGFsaWNBbmdsZSAwIC9MZWFkaW5nIDE0'
    'OCAvTWF4V2lkdGggMTAwMCAvTWlzc2luZ1dpZHRoIDUwMCAvU3RlbUggOTEgL1N0ZW1WIDU4IAogIC9UeXBlIC9Gb250RGVzY3Jp'
    'cHRvciAvWEhlaWdodCA1NTMKPj4gL1N1YnR5cGUgL0NJREZvbnRUeXBlMCAvVHlwZSAvRm9udCAKICAvVyBbIDEgWyAyMDcgMjcw'
    'IDM0MiA0NjcgNDYyIDc5NyA3MTAgMjM5IDM3NCBdIDEwIFsgMzc0IDQyMyA2MDUgMjM4IDM3NSAyMzggMzM0IDQ2MiBdIDE4IDI2'
    'IDQ2MiAyNyAyOCAyMzggCiAgMjkgMzEgNjA1IDMyIFsgMzQ0IDc0OCA2ODQgNTYwIDY5NSA3MzkgNTYzIDUxMSA3MjkgNzkzIAog'
    'IDMxOCAzMTIgNjY2IDUyNiA4OTYgNzU4IDc3MiA1NDQgNzcyIDYyOCAKICA0NjUgNjA3IDc1MyA3MTEgOTcyIDY0NyA2MjAgNjA3'
    'IDM3NCAzMzMgCiAgMzc0IDYwNiA1MDAgMjM5IDQxNyA1MDMgNDI3IDUyOSA0MTUgMjY0IAogIDQ0NCA1MTggMjQxIDIzMCA0OTUg'
    'MjI4IDc5MyA1MjcgNTI0IF0gODEgWyA1MjQgNTA0IDMzOCAzMzYgMjc3IDUxNyA0NTAgNjUyIDQ2NiA0NTIgCiAgNDA3IDM3MCAy'
    'NTggMzcwIDYwNSBdIF0KPj4gXSAvRW5jb2RpbmcgL1VuaUdCLVVDUzItSCAvTmFtZSAvRjIgL1N1YnR5cGUgL1R5cGUwIC9UeXBl'
    'IC9Gb250Cj4+CmVuZG9iago0IDAgb2JqCjw8Ci9Db250ZW50cyA4IDAgUiAvTWVkaWFCb3ggWyAwIDAgNTk1LjI3NTYgODQxLjg4'
    'OTggXSAvUGFyZW50IDcgMCBSIC9SZXNvdXJjZXMgPDwKL0ZvbnQgMSAwIFIgL1Byb2NTZXQgWyAvUERGIC9UZXh0IC9JbWFnZUIg'
    'L0ltYWdlQyAvSW1hZ2VJIF0KPj4gL1JvdGF0ZSAwIC9UcmFucyA8PAoKPj4gCiAgL1R5cGUgL1BhZ2UKPj4KZW5kb2JqCjUgMCBv'
    'YmoKPDwKL1BhZ2VNb2RlIC9Vc2VOb25lIC9QYWdlcyA3IDAgUiAvVHlwZSAvQ2F0YWxvZwo+PgplbmRvYmoKNiAwIG9iago8PAov'
    'QXV0aG9yIChhbm9ueW1vdXMpIC9DcmVhdGlvbkRhdGUgKEQ6MjAyNjEwMTAwMTI3MjktMDcnMDAnKSAvQ3JlYXRvciAoYW5vbnlt'
    'b3VzKSAvS2V5d29yZHMgKCkgL01vZERhdGUgKEQ6MjAyNjEwMTAwMTI3MjktMDcnMDAnKSAvUHJvZHVjZXIgKFJlcG9ydExhYiBQ'
    'REYgTGlicmFyeSAtIFwob3BlbnNvdXJjZVwpKSAKICAvU3ViamVjdCAodW5zcGVjaWZpZWQpIC9UaXRsZSAodW50aXRsZWQpIC9U'
    'cmFwcGVkIC9GYWxzZQo+PgplbmRvYmoKNyAwIG9iago8PAovQ291bnQgMSAvS2lkcyBbIDQgMCBSIF0gL1R5cGUgL1BhZ2VzCj4+'
    'CmVuZG9iago4IDAgb2JqCjw8Ci9GaWx0ZXIgWyAvQVNDSUk4NURlY29kZSAvRmxhdGVEZWNvZGUgXSAvTGVuZ3RoIDc3OQo+Pgpz'
    'dHJlYW0KR2FzYW4/I3BXRSY6RGctPTw9UDNYLEstSD9QOSE8ZUFOUE8pQGRBNDhzJU8nS2NHQSgma2pkaWgza3I8Oi9cL2I7MDla'
    'MW9ea2lPIVhNQC4jXkRhRm5MYlFbNShXYktrdWBJImUwNyRqKzRAaUFOV0w+blgkYHU3VHFqQT9cPkJzU144KDg3LE1KP2RRY2RR'
    'OT4zMiNqP0BxZXRtPmhSSWMzUlwwLnI0YUI4Il4yR0FKSTxEa3RJbStHNGArIWRkSiElXi5uUnFIcUNFQmpBUDM8UFFCay01OFdN'
    'Nyh1YGc7VW1MLkhdXztiPExCaVtiXyFDMmZXcWs3OiQnM3AiKyxqP2QjXVM4a21aZ1pCNi9uL2FpUlAmIyU4L2Flb2lvJDRAc0dP'
    'RU1wSCssM1svTjNyVSYpalFSPDtMbyhoMDNeMUkpaC04ZEZLOGpfMCFja1VLXEBNL0UtPFlMSFxrUmllO3BIMW10TjE6JHBZS0sy'
    'cCVUUUEuZC0sXXIqSy5xOUUvcE0jQGUjLlxjK2tQcFNtI0dULz82R15VO1FKY2lPIm8nYjRTKTlfcGQ7dVlEMCg6W20oXz8pY0Vn'
    'PFhRI0J0UThbYGhsWDdbRmRVbnJRWnFWTHNMcjBZQG1HU3A3JW48WkFwVFJFb2dwREcjInJcOk46OjtcVF0kOVRwQS9bZiIncVpk'
    'XHBMLjQlPmFnb28/Ky9NWVs4Z2tQYDw/akE6NyI+LEU/dXQyaDZVcmpJdC0jdGdvO1BcQEU4ZVZnJjxqb1gqWG5MSGtmNURcJmNG'
    'Xj9wSW5SWEIpKz4zQGk2SHJSNC9FLz5cMS5rR1tTREsvbC47Xj9HX1cyOiYtSUt0TD03WyFCM0wyK0NuJChXK2R0aSMiNThNb2tD'
    'S0FYK2tDbVYyZ0EzbEkoTVFFY1s7QEc2QyhnLGowbzRjQylCcEs/XWUkQ25GOycrV1ZPcyIiVEI0ViFCRG4oO2g2RmA+MCpXVG9v'
    'KGNtVS5URCNGc1djMUZBUzg+NjMmLjJSPlwhT1tIUTJafj5lbmRzdHJlYW0KZW5kb2JqCnhyZWYKMCA5CjAwMDAwMDAwMDAgNjU1'
    'MzUgZiAKMDAwMDAwMDA2MSAwMDAwMCBuIAowMDAwMDAwMTAyIDAwMDAwIG4gCjAwMDAwMDAyMDkgMDAwMDAgbiAKMDAwMDAwMTE0'
    'MiAwMDAwMCBuIAowMDAwMDAxMzQ1IDAwMDAwIG4gCjAwMDAwMDE0MTMgMDAwMDAgbiAKMDAwMDAwMTY3NCAwMDAwMCBuIAowMDAw'
    'MDAxNzMzIDAwMDAwIG4gCnRyYWlsZXIKPDwKL0lEIApbPDU5ZjdmZDcyNDFjNzI2YjgzNjYwYjZjZDllNDUwNGQxPjw1OWY3ZmQ3'
    'MjQxYzcyNmI4MzY2MGI2Y2Q5ZTQ1MDRkMT5dCiUgUmVwb3J0TGFiIGdlbmVyYXRlZCBQREYgZG9jdW1lbnQgLS0gZGlnZXN0IChv'
    'cGVuc291cmNlKQoKL0luZm8gNiAwIFIKL1Jvb3QgNSAwIFIKL1NpemUgOQo+PgpzdGFydHhyZWYKMjYwMgolJUVPRgo='
)


def fixture_bytes(ext):
    if ext == 'txt':
        return FICTIONAL_TEXT.encode('utf-8')
    if ext == 'pdf':
        return base64.b64decode(_PDF_BASE64, validate=True)
    if ext != 'docx':
        raise ValueError('unsupported fictional fixture')
    from docx import Document
    document = Document()
    document.core_properties.author = 'ApplyPilot fictional regression'
    document.core_properties.created = datetime(2000, 1, 1, tzinfo=timezone.utc)
    document.core_properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    document.add_heading('虚构简历文件解析测试', 0)
    for line in FICTIONAL_TEXT.splitlines():
        document.add_paragraph(line)
    raw = io.BytesIO()
    document.save(raw)
    result = io.BytesIO()
    with zipfile.ZipFile(raw) as original, zipfile.ZipFile(result, 'w') as normalized:
        for name in sorted(original.namelist()):
            entry = zipfile.ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            normalized.writestr(entry, original.read(name), compress_type=zipfile.ZIP_DEFLATED)
    return result.getvalue()


if __name__ == '__main__':
    import json
    print(json.dumps({ext: base64.b64encode(fixture_bytes(ext)).decode() for ext in ('txt', 'pdf', 'docx')}))
